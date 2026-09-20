"""レポート生成の回帰テスト。

修正前は reviews_count が文字列だと ValueError: Cannot specify ',' with 's' で
Word 生成が丸ごと落ちていた (LLM は「350件」等を返しうる)。
"""

import copy

import pytest
from conftest import SAMPLE_DATA
from docx import Document

import report_engine
from report_engine import as_int, generate_full_report_pack


@pytest.fixture(autouse=True)
def _no_network(monkeypatch):
    """地図タイル取得と住所ジオコーディングを止め、テストを外部依存なしにする"""
    monkeypatch.setattr(report_engine, "geocode_address", lambda _addr: None)
    monkeypatch.setattr(report_engine, "_fetch_gsi_tile_safe", lambda *a, **k: None, raising=False)

    def _no_tiles(*args, **kwargs):
        raise OSError("network disabled in tests")

    monkeypatch.setattr(report_engine.urllib.request, "urlopen", _no_tiles)


@pytest.mark.parametrize(
    ("raw", "expected"),
    [(350, 350), ("350件", 350), ("約 1,200", 1200), ("", 0), (None, 0), (12.7, 12), (True, 0), ("不明", 0)],
)
def test_as_int_normalises_llm_values(raw, expected):
    assert as_int(raw) == expected


def test_string_reviews_count_does_not_crash(tmp_path):
    data = copy.deepcopy(SAMPLE_DATA)
    data["spots"][0]["reviews_count"] = "350件"
    data["spots"][1]["reviews_count"] = None
    pack = generate_full_report_pack(data, tmp_path)
    assert pack["docx_path"].exists()


def test_generates_every_artifact(tmp_path):
    pack = generate_full_report_pack(copy.deepcopy(SAMPLE_DATA), tmp_path)
    for key in ("docx_path", "mobile_docx_path", "csv_path", "map_path"):
        assert pack[key] is not None and pack[key].exists(), f"{key} が生成されていない"
    assert len(list(tmp_path.glob("*.docx"))) == 2, "互換用の重複 docx が復活している"


def test_spots_without_photos_render_cleanly(tmp_path):
    """写真が無いとき、空の写真枠や『写真準備中』を出さないこと"""
    pack = generate_full_report_pack(copy.deepcopy(SAMPLE_DATA), tmp_path)
    doc = Document(str(pack["docx_path"]))
    assert "写真準備中" not in "\n".join(p.text for p in doc.paragraphs)
    # 写真テーブル(2行×N列)が無いので、表は諸元テーブルとコールアウトのみ
    assert all(len(t.rows) != 2 or len(t.columns) != 2 or t.cell(0, 0).text for t in doc.tables)


def test_sources_are_cited_in_report(tmp_path):
    pack = generate_full_report_pack(copy.deepcopy(SAMPLE_DATA), tmp_path)
    text = "\n".join(p.text for p in Document(str(pack["docx_path"])).paragraphs)
    assert "調査の出典" in text
    assert "銀座 個室 鮨" in text


def test_key_topics_are_rendered(tmp_path):
    """CSV にしか出ていなかった特徴キーワードが Word にも載ること"""
    pack = generate_full_report_pack(copy.deepcopy(SAMPLE_DATA), tmp_path)
    for path in (pack["docx_path"], pack["mobile_docx_path"]):
        doc = Document(str(path))
        blob = "\n".join(p.text for p in doc.paragraphs)
        blob += "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
        assert "個室あり" in blob, f"{path.name} に key_topics が出ていない"


def test_review_heading_matches_actual_count(tmp_path):
    """『厳選8件』と書いておいて実際は少ない、という食い違いを作らないこと"""
    data = copy.deepcopy(SAMPLE_DATA)
    data["spots"][0]["reviews"] = ["口コミA", "口コミB"]
    pack = generate_full_report_pack(data, tmp_path)
    text = "\n".join(p.text for p in Document(str(pack["mobile_docx_path"])).paragraphs)
    assert "（2件）" in text
    assert "8件" not in text


def test_csv_has_header_and_one_row_per_spot(tmp_path):
    pack = generate_full_report_pack(copy.deepcopy(SAMPLE_DATA), tmp_path)
    lines = pack["csv_path"].read_text(encoding="utf-8-sig").strip().splitlines()
    assert len(lines) == 1 + len(SAMPLE_DATA["spots"])
